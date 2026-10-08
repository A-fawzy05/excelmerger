import os, re, sys
from contextlib import contextmanager
from copy import copy
from itertools import islice

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from PySide6.QtCore import QEvent, QPoint, QPointF, Qt, Signal
from PySide6.QtGui import QColor, QIcon, QPainter, QPolygonF
from PySide6.QtWidgets import (
    QApplication, QDialog, QFileDialog, QFrame, QHBoxLayout, QLabel, QListWidget, QListWidgetItem,
    QMainWindow, QMessageBox, QPushButton, QStackedWidget, QTableWidget, QTableWidgetItem,
    QVBoxLayout, QWidget,
)

APP_NAME = "Excel Editor made by Amr"
SOURCE_COL = "اسم_الملف"
PREVIEW_ROWS = 200
ALIASES = {  # canonical name: other spellings seen in the files
    "الاسم": ["الإسم", "اسم"],
    "الهاتف": ["رقم الهاتف", "موبايل"],
}
DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "0123456789" * 2)  # Arabic/Persian digits -> 0-9


def resource(name):
    return os.path.join(getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__))), name)

# ---------------------------------------------------------------- data logic

def norm(s):  # headers only: drop tashkeel/tatweel, unify Arabic letter variants
    s = re.sub(r"[\u064B-\u065F\u0640]", "", str(s))
    s = re.sub("[ىی]", "ي", s).replace("ک", "ك")
    s = re.sub("[أإآ]", "ا", s)
    return re.sub(r"\s+", " ", s).strip()

LOOKUP = {norm(a): k for k, v in ALIASES.items() for a in [k, *v]}

def kval(v):  # key value: 123, "123", "١٢٣", 123.0 all match
    s = str(v).strip().translate(DIGITS)
    return s[:-2] if s.endswith(".0") else s

def sheet_headers(p):  # {sheet name: [canonical header names]}
    wb = load_workbook(p, read_only=True)
    out = {}
    for ws in wb.worksheets:
        row = next(ws.iter_rows(max_row=1, values_only=True), ())
        out[ws.title] = [LOOKUP.get(norm(v), norm(v)) for v in row if v is not None]
    wb.close()
    return out

def clone(src, dst):  # value + look of one cell
    dst.value = src.value
    if src.has_style:
        dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
        dst.alignment, dst.protection = copy(src.alignment), copy(src.protection)
        dst.number_format = src.number_format

def build(paths, sheets, keys=()):
    cols, hdr, width, recs = [], {}, {}, {}
    for p in paths:
        wb = load_workbook(p, data_only=True)
        for ws in wb.worksheets:
            if ws.title not in sheets:
                continue
            tag = f"{os.path.basename(p)} [{ws.title}]"
            it = ws.iter_rows()
            head = next(it, None)
            if not head:
                continue
            names = [LOOKUP.get(norm(c.value), norm(c.value)) if c.value is not None else "" for c in head]
            for n, c in zip(names, head):
                if n and n not in hdr:  # first sheet that has the column sets its header style and width
                    cols.append(n)
                    hdr[n] = c
                    width[n] = ws.column_dimensions[c.column_letter].width
            for r in it:
                if all(c.value is None for c in r):
                    continue
                d = {}
                for n, c in zip(names, r):
                    if n and n not in d:  # ponytail: duplicate header in one sheet keeps first only
                        d[n] = c
                if keys and all(k in d and d[k].value not in (None, "") for k in keys):
                    key = tuple(kval(d[k].value) for k in keys)
                else:
                    key = object()  # no keys chosen, or key cell empty: row stays separate
                rec = recs.setdefault(key, {"cells": {}, "files": [], "h": ws.row_dimensions[r[0].row].height})
                rec["files"].append(tag)
                for n, c in d.items():
                    cur = rec["cells"].get(n)
                    if cur is None or cur.value in (None, ""):  # ponytail: on conflict first source wins
                        rec["cells"][n] = c
    if not cols:
        raise ValueError("لا توجد بيانات في الأوراق المختارة")
    missing = [k for k in keys if k not in cols]
    if missing:
        raise ValueError("عمود المفتاح غير موجود: " + ", ".join(missing))
    cols = list(keys) + [c for c in cols if c not in keys]  # key columns first
    return cols, hdr, width, recs

def write(out, cols, hdr, width, rows):  # rows: iterable of records from build()
    wb = Workbook()
    ws = wb.active
    ws.sheet_view.rightToLeft = True
    last = len(cols) + 1
    for j, n in enumerate(cols, 1):
        clone(hdr[n], ws.cell(1, j))
        ws.cell(1, j).value = n  # canonical name, not the original spelling
        if width[n]:
            ws.column_dimensions[get_column_letter(j)].width = width[n]
    clone(hdr[cols[0]], ws.cell(1, last))
    ws.cell(1, last).value = SOURCE_COL
    ws.column_dimensions[get_column_letter(last)].width = 30
    for i, rec in enumerate(rows, 2):
        for j, n in enumerate(cols, 1):
            if n in rec["cells"]:
                clone(rec["cells"][n], ws.cell(i, j))
        ws.cell(i, last, "، ".join(dict.fromkeys(rec["files"])))
        if rec["h"]:
            ws.row_dimensions[i].height = rec["h"]
    wb.save(out)

def split_groups(recs, col):  # {file name: [records]}, one group per distinct value of col
    groups = {}
    for rec in recs.values():
        v = getattr(rec["cells"].get(col), "value", None)
        groups.setdefault("" if v in (None, "") else kval(v), []).append(rec)
    names, used = {}, {}
    for k in groups:
        base = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", k or "فارغ").strip(" .")[:80] or "فارغ"
        n = used[base.lower()] = used.get(base.lower(), 0) + 1  # two values that clean to one name get _2, _3
        names[base + (f"_{n}" if n > 1 else "") + ".xlsx"] = groups[k]
    return names

# ------------------------------------------------------------------ palette

P = {  # deep emerald gradient backdrop, white panels, Excel-green accent
    "BG": "#F2F4F6", "SURFACE": "#FFFFFF", "BORDER": "#D5DCE3",
    "TEXT": "#16202A", "MUTED": "#5A6773", "ON_DARK": "#D8F0E4",
    "DEEP": "#082A21", "MID": "#0D5640", "GLOW": "#15855E",
    "ACCENT_HOVER": "#0B6340", "ACCENT_SOFT": "#E2F1E9", "ACCENT": "#0E7A4F",  # longest names first for replace
}

QSS = """
* { font-family: "Segoe UI", Tahoma; font-size: 14px; color: @TEXT; }
QMainWindow { background: @DEEP; }
QDialog { background: @BG; }
#page { background: qlineargradient(x1:0, y1:0, x2:1, y2:1, stop:0 @DEEP, stop:0.55 @MID, stop:1 @GLOW); }
#panel { background: @SURFACE; border: 1px solid @BORDER; border-radius: 14px; }
QLabel { background: transparent; }
QLabel#h1 { font-size: 30px; font-weight: 700; color: #FFFFFF; }
QLabel#h2 { font-size: 22px; font-weight: 700; color: #FFFFFF; }
QLabel#lead { font-size: 15px; color: @ON_DARK; }
QLabel#sub { color: @MUTED; }
QLabel#field { font-weight: 600; }
QLabel#hint { color: @MUTED; font-size: 12px; }
QLabel#cardtitle { font-size: 19px; font-weight: 700; }
QPushButton { background: @SURFACE; border: 1px solid @BORDER; border-radius: 8px; padding: 9px 20px; }
QPushButton:hover, QPushButton:focus { background: @ACCENT_SOFT; border-color: @ACCENT; }
QPushButton#primary { background: #FFFFFF; border: 1px solid #FFFFFF; color: @DEEP; font-weight: 700; }
QPushButton#primary:hover, QPushButton#primary:focus { background: @ACCENT_SOFT; border-color: @ACCENT_SOFT; }
QPushButton#ghost { background: transparent; border: 1px solid rgba(255, 255, 255, 150); color: #FFFFFF; font-weight: 600; }
QPushButton#ghost:hover, QPushButton#ghost:focus { background: rgba(255, 255, 255, 35); border-color: #FFFFFF; }
QPushButton#link { background: transparent; border: none; color: @ON_DARK; font-weight: 600; padding: 6px 4px; }
QPushButton#link:hover { background: transparent; color: #FFFFFF; text-decoration: underline; }
QPushButton#card { background: @SURFACE; border: 1px solid @BORDER; border-radius: 14px; }
QPushButton#card:hover, QPushButton#card:focus { background: @ACCENT_SOFT; border: 2px solid @ACCENT; }
QPushButton#drop { background: @SURFACE; border: 1px solid @BORDER; border-radius: 8px; padding: 9px 12px 9px 34px; text-align: left; }
QPushButton#drop:hover, QPushButton#drop:focus { background: @SURFACE; border-color: @ACCENT; }
QPushButton#drop[empty="true"] { color: @MUTED; }
#popup { background: @SURFACE; border: 1px solid @ACCENT; border-radius: 8px; }
#popup QListWidget { border: none; background: transparent; padding: 4px; }
QListWidget, QTableWidget { background: @SURFACE; border: 1px solid @BORDER; border-radius: 8px; padding: 5px 8px; }
QListWidget:focus { border-color: @ACCENT; }
QListWidget { outline: 0; }
QListWidget::item { padding: 5px 4px; border-radius: 4px; }
QListWidget::item:hover, QListWidget::item:selected { background: @ACCENT_SOFT; color: @TEXT; }
QAbstractItemView::indicator { width: 16px; height: 16px; border: 1.5px solid @MUTED; border-radius: 4px; background: @SURFACE; }
QAbstractItemView::indicator:checked { background: @ACCENT; border-color: @ACCENT; }
QHeaderView::section { background: @BG; border: none; border-bottom: 1px solid @BORDER; padding: 6px; font-weight: 600; }
QTableWidget { gridline-color: @BORDER; alternate-background-color: #F8FAFB; }
"""

def make_qss():
    out = QSS
    for k, v in P.items():
        out = out.replace("@" + k, v)
    return out

# ------------------------------------------------------------------ widgets

@contextmanager
def busy():
    QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
    try:
        yield
    finally:
        QApplication.restoreOverrideCursor()

def box(parent, icon, title, text, ok="موافق", cancel=None):  # Arabic buttons, returns True on ok
    m = QMessageBox(icon, title, text, parent=parent)
    yes = m.addButton(ok, QMessageBox.ButtonRole.AcceptRole)
    if cancel:
        m.addButton(cancel, QMessageBox.ButtonRole.RejectRole)
    m.exec()
    return m.clickedButton() is yes

class CheckCombo(QPushButton):
    """Drop-down button. Click opens a popup list: checkboxes (stays open) or, with single=True, pick one (closes)."""
    changed = Signal()

    def __init__(self, placeholder, single=False):
        super().__init__()
        self.ph, self.single = placeholder, single
        self.setObjectName("drop")
        self.setCursor(Qt.CursorShape.PointingHandCursor)
        self.pop = QFrame(self, Qt.WindowType.Popup)  # closes by itself on any click outside
        self.pop.setObjectName("popup")
        self.list = QListWidget()
        v = QVBoxLayout(self.pop)
        v.setContentsMargins(0, 0, 0, 0)
        v.addWidget(self.list)
        self.list.viewport().installEventFilter(self)
        self.clicked.connect(self.open)
        self._sync()

    def open(self):
        n = self.list.count()
        if not n:
            return
        self.pop.setFixedSize(self.width(), min(n, 8) * self.list.sizeHintForRow(0) + 20)
        self.pop.move(self.mapToGlobal(QPoint(0, self.height() + 2)))
        self.pop.show()

    def eventFilter(self, obj, ev):
        if obj is self.list.viewport() and ev.type() == QEvent.Type.MouseButtonRelease:
            it = self.list.itemAt(ev.position().toPoint())
            if it:
                self._toggle(it)
            return True  # we handle the click ourselves, so the checkbox never double-toggles
        return super().eventFilter(obj, ev)

    def _toggle(self, it):
        on = it.checkState() == Qt.CheckState.Checked
        if self.single:
            for i in range(self.list.count()):
                self.list.item(i).setCheckState(Qt.CheckState.Unchecked)
            self.pop.hide()
        it.setCheckState(Qt.CheckState.Unchecked if on else Qt.CheckState.Checked)
        self._sync()

    def checked(self):
        items = [self.list.item(i) for i in range(self.list.count())]
        return [it.text() for it in items if it.checkState() == Qt.CheckState.Checked]

    def value(self):  # single mode: the chosen name or ""
        return (self.checked() or [""])[0]

    def set_items(self, names, on):  # on: names to tick
        self.list.clear()
        for n in names:
            it = QListWidgetItem(n)
            it.setFlags(Qt.ItemFlag.ItemIsUserCheckable | Qt.ItemFlag.ItemIsEnabled)
            it.setCheckState(Qt.CheckState.Checked if n in on else Qt.CheckState.Unchecked)
            self.list.addItem(it)
        self._sync()

    def _sync(self):
        c = self.checked()
        self.setText(("، ".join(c) if len(c) <= 2 else f"{len(c)} محدد") if c else self.ph)
        self.setProperty("empty", not c)
        self.style().unpolish(self)
        self.style().polish(self)
        self.changed.emit()

    def paintEvent(self, e):
        super().paintEvent(e)
        p = QPainter(self)
        p.setRenderHint(QPainter.RenderHint.Antialiasing)
        p.setPen(Qt.PenStyle.NoPen)
        p.setBrush(QColor(P["MUTED"]))
        y = self.height() / 2
        p.drawPolygon(QPolygonF([QPointF(12, y - 3), QPointF(22, y - 3), QPointF(17, y + 3)]))
        p.end()

class Sources(QWidget):
    """Shared by both pages: file list, sheet drop-down, key-column drop-down."""
    changed = Signal()

    def __init__(self, keys_placeholder):
        super().__init__()
        self.paths, self.info, self.known, self.cols = [], {}, set(), []
        self.files = QListWidget()
        self.files.setFixedHeight(104)
        self.sheets = CheckCombo("اختر الأوراق")
        self.keys = CheckCombo(keys_placeholder)
        self.sheets.changed.connect(self.refresh_cols)
        add, clear = QPushButton("إضافة ملفات"), QPushButton("مسح القائمة")
        add.clicked.connect(self.add)
        clear.clicked.connect(self.clear)
        bar = QHBoxLayout()
        bar.addWidget(add)
        bar.addWidget(clear)
        bar.addStretch()
        self.filebox = QWidget()  # list + buttons, placed on the page by Page.field()
        v = QVBoxLayout(self.filebox)
        v.setContentsMargins(0, 0, 0, 0)
        v.addWidget(self.files)
        v.addLayout(bar)

    def add(self):
        paths, _ = QFileDialog.getOpenFileNames(self.window(), "اختر ملفات Excel", "", "Excel (*.xlsx)")
        self.load(paths)

    def load(self, paths):
        for p in paths:
            if p in self.paths:
                continue
            try:
                self.info[p] = sheet_headers(p)
            except Exception as e:
                box(self.window(), QMessageBox.Icon.Critical, "خطأ", f"{os.path.basename(p)}: {e}")
                continue
            self.paths.append(p)
            self.files.addItem(os.path.basename(p))
        names = list(dict.fromkeys(s for p in self.paths for s in self.info[p]))
        on = set(self.sheets.checked())
        self.sheets.set_items(names, {n for n in names if n in on or n not in self.known})  # new sheets start ticked
        self.known |= set(names)

    def clear(self):
        self.paths.clear(); self.info.clear(); self.known.clear(); self.files.clear()
        self.sheets.set_items([], set())

    def refresh_cols(self):
        ch, on = set(self.sheets.checked()), set(self.keys.checked())
        self.cols = list(dict.fromkeys(c for p in self.paths for s, h in self.info[p].items() if s in ch for c in h))
        self.keys.set_items(self.cols, on)
        self.changed.emit()

    def args(self):
        if not self.paths:
            box(self.window(), QMessageBox.Icon.Warning, "تنبيه", "اختر الملفات أولاً")
        elif not self.sheets.checked():
            box(self.window(), QMessageBox.Icon.Warning, "تنبيه", "اختر ورقة واحدة على الأقل")
        else:
            return self.paths, set(self.sheets.checked()), self.keys.checked()

def show_preview(parent, cols, recs):
    d = QDialog(parent)
    d.setWindowTitle(f"معاينة: {len(recs)} صف، {len(cols) + 1} عمود (عرض أول {PREVIEW_ROWS})")
    d.resize(1000, 540)
    heads = cols + [SOURCE_COL]
    rows = list(islice(recs.values(), PREVIEW_ROWS))
    t = QTableWidget(len(rows), len(heads))
    t.setHorizontalHeaderLabels(heads)
    t.setAlternatingRowColors(True)
    t.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
    for i, rec in enumerate(rows):
        for j, c in enumerate(cols):
            v = getattr(rec["cells"].get(c), "value", None)
            t.setItem(i, j, QTableWidgetItem("" if v is None else str(v)))
        t.setItem(i, len(cols), QTableWidgetItem("، ".join(dict.fromkeys(rec["files"]))))
    t.resizeColumnsToContents()
    for j in range(len(heads)):
        t.setColumnWidth(j, min(t.columnWidth(j), 260))
    QVBoxLayout(d).addWidget(t)
    d.exec()

# -------------------------------------------------------------------- pages

class Page(QWidget):
    def __init__(self, title, back):
        super().__init__()
        self.setObjectName("page")
        self.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)  # lets the gradient paint on a plain QWidget
        root = QVBoxLayout(self)
        root.setContentsMargins(32, 24, 32, 24)
        root.setSpacing(16)
        top = QHBoxLayout()
        b = QPushButton("رجوع")
        b.setObjectName("link")
        b.setCursor(Qt.CursorShape.PointingHandCursor)
        b.clicked.connect(back)
        t = QLabel(title)
        t.setObjectName("h2")
        top.addWidget(b)
        top.addSpacing(12)
        top.addWidget(t)
        top.addStretch()
        self.panel = QFrame()
        self.panel.setObjectName("panel")
        self.form = QVBoxLayout(self.panel)
        self.form.setContentsMargins(24, 20, 24, 20)
        self.form.setSpacing(6)
        self.bar = QHBoxLayout()
        root.addLayout(top)
        root.addWidget(self.panel)
        root.addStretch()
        root.addLayout(self.bar)

    def field(self, label, widget, hint=None):
        l = QLabel(label)
        l.setObjectName("field")
        self.form.addSpacing(8)
        self.form.addWidget(l)
        self.form.addWidget(widget)
        if hint:
            h = QLabel(hint)
            h.setObjectName("hint")
            h.setWordWrap(True)
            self.form.addWidget(h)

    def actions(self, *buttons):  # first button ends up at the right edge (primary)
        for b in buttons:
            b.setCursor(Qt.CursorShape.PointingHandCursor)
            self.bar.addWidget(b)
        self.bar.addStretch()

class MergePage(Page):
    def __init__(self, back):
        super().__init__("دمج الملفات", back)
        self.src = Sources("اختياري")
        self.field("الملفات", self.src.filebox)
        self.field("الأوراق", self.src.sheets, "اختر ورقة أو أكثر. الأوراق التي تحمل نفس الاسم تُدمج من كل الملفات.")
        self.field("أعمدة الربط", self.src.keys,
                   "اختياري. الصفوف التي تحمل نفس القيمة في هذه الأعمدة تُدمج في صف واحد.")
        go, view = QPushButton("دمج وحفظ"), QPushButton("معاينة")
        go.setObjectName("primary")
        view.setObjectName("ghost")
        go.clicked.connect(self.run)
        view.clicked.connect(self.preview)
        self.actions(go, view)

    def preview(self):
        a = self.src.args()
        if not a:
            return
        try:
            with busy():
                cols, _, _, recs = build(*a)
        except Exception as e:
            return box(self.window(), QMessageBox.Icon.Critical, "خطأ", str(e))
        show_preview(self.window(), cols, recs)

    def run(self):
        a = self.src.args()
        if not a:
            return
        out, _ = QFileDialog.getSaveFileName(self.window(), "حفظ الملف",
                                             os.path.join(os.path.dirname(self.src.paths[0]), "merged.xlsx"),
                                             "Excel (*.xlsx)")
        if not out:
            return
        try:
            with busy():
                cols, hdr, width, recs = build(*a)
                write(out, cols, hdr, width, recs.values())
        except Exception as e:
            return box(self.window(), QMessageBox.Icon.Critical, "خطأ", str(e))
        box(self.window(), QMessageBox.Icon.Information, "تم", f"الصفوف: {len(recs)}\nالأعمدة: {len(cols) + 1}\n\n{out}")

class SplitPage(Page):
    def __init__(self, back):
        super().__init__("تقسيم إلى ملفات", back)
        self.src = Sources("اختياري")
        self.col = CheckCombo("اختر عمود التقسيم", single=True)
        self.src.changed.connect(self.refresh_col)
        self.field("الملفات", self.src.filebox)
        self.field("الأوراق", self.src.sheets)
        self.field("عمود التقسيم", self.col, "يُنشأ ملف لكل قيمة مختلفة في هذا العمود.")
        go, view = QPushButton("تقسيم وحفظ"), QPushButton("معاينة")
        go.setObjectName("primary")
        view.setObjectName("ghost")
        go.clicked.connect(self.run)
        view.clicked.connect(self.preview)
        self.actions(go, view)

    def refresh_col(self):
        self.col.set_items(self.src.cols, set(self.col.checked()))

    def plan(self):  # (cols, hdr, width, files) or None after showing why
        a = self.src.args()
        if not a:
            return
        col = self.col.value()
        if not col:
            return box(self.window(), QMessageBox.Icon.Warning, "تنبيه", "اختر عمود التقسيم")
        try:
            with busy():
                cols, hdr, width, recs = build(*a)
                if col not in cols:
                    raise ValueError("عمود التقسيم غير موجود: " + col)
                return cols, hdr, width, split_groups(recs, col)
        except Exception as e:
            return box(self.window(), QMessageBox.Icon.Critical, "خطأ", str(e))

    def preview(self):  # shows the files that would be created: file name + row count
        p = self.plan()
        if not p:
            return
        files = p[3]
        d = QDialog(self.window())
        d.setWindowTitle(f"معاينة التقسيم: {len(files)} ملف")
        d.resize(520, 480)
        t = QTableWidget(len(files), 2)
        t.setHorizontalHeaderLabels(["اسم الملف", "عدد الصفوف"])
        t.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        t.setAlternatingRowColors(True)
        for i, (f, rows) in enumerate(files.items()):
            t.setItem(i, 0, QTableWidgetItem(f))
            t.setItem(i, 1, QTableWidgetItem(str(len(rows))))
        t.setColumnWidth(0, 330)
        QVBoxLayout(d).addWidget(t)
        d.exec()

    def run(self):
        p = self.plan()
        if not p:
            return
        cols, hdr, width, files = p
        folder = QFileDialog.getExistingDirectory(self.window(), "مجلد حفظ الملفات", os.path.dirname(self.src.paths[0]))
        if not folder:
            return
        clash = sum(os.path.exists(os.path.join(folder, f)) for f in files)
        msg = f"سيتم إنشاء {len(files)} ملف (حسب عمود: {self.col.value()}) في:\n{folder}"
        if clash:
            msg += f"\n\nتنبيه: {clash} ملف موجود وسيتم استبداله"
        if not box(self.window(), QMessageBox.Icon.Question, "تأكيد", msg, "متابعة", "إلغاء"):
            return
        try:
            with busy():
                for f, rows in files.items():
                    write(os.path.join(folder, f), cols, hdr, width, rows)
        except Exception as e:
            return box(self.window(), QMessageBox.Icon.Critical, "خطأ", str(e))
        box(self.window(), QMessageBox.Icon.Information, "تم", f"تم إنشاء {len(files)} ملف في:\n{folder}")

class Home(QWidget):
    def __init__(self, go):
        super().__init__()
        self.setObjectName("page")
        self.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        root = QVBoxLayout(self)
        root.setContentsMargins(48, 40, 48, 40)
        title = QLabel(APP_NAME)
        title.setObjectName("h1")
        title.setAlignment(Qt.AlignmentFlag.AlignHCenter)
        sub = QLabel("اختر العملية التي تريدها")
        sub.setObjectName("lead")
        sub.setAlignment(Qt.AlignmentFlag.AlignHCenter)
        row = QHBoxLayout()
        row.setSpacing(20)
        row.addWidget(self.card("دمج الملفات", "اجمع عدة ملفات وأوراق Excel في ملف واحد، مع إمكانية الربط بعمود مفتاح.", lambda: go(1)))
        row.addWidget(self.card("تقسيم إلى ملفات", "قسّم البيانات إلى ملفات منفصلة، ملف لكل قيمة في العمود الذي تختاره.", lambda: go(2)))
        root.addStretch(2)
        root.addWidget(title)
        root.addWidget(sub)
        root.addSpacing(28)
        root.addLayout(row)
        root.addStretch(3)

    def card(self, head, body, fn):
        b = QPushButton()
        b.setObjectName("card")
        b.setCursor(Qt.CursorShape.PointingHandCursor)
        b.setMinimumHeight(170)
        b.clicked.connect(fn)
        lay = QVBoxLayout(b)
        lay.setContentsMargins(24, 22, 24, 22)
        t, d = QLabel(head), QLabel(body)
        t.setObjectName("cardtitle")
        d.setObjectName("sub")
        d.setWordWrap(True)
        for w in (t, d):
            w.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        lay.addWidget(t)
        lay.addWidget(d)
        lay.addStretch()
        return b

class Main(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_NAME)
        self.resize(940, 780)
        self.setMinimumSize(780, 640)
        self.stack = QStackedWidget()
        self.setCentralWidget(self.stack)
        go = self.stack.setCurrentIndex
        self.stack.addWidget(Home(go))
        self.stack.addWidget(MergePage(lambda: go(0)))
        self.stack.addWidget(SplitPage(lambda: go(0)))

if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setWindowIcon(QIcon(resource("icon.png")))
    app.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
    app.setStyleSheet(make_qss())
    win = Main()
    win.setFocusPolicy(Qt.FocusPolicy.ClickFocus)
    win.show()
    win.setFocus()  # no card starts highlighted; Tab still reaches them
    sys.exit(app.exec())